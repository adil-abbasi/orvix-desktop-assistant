from app.session_memory import get_last_project, set_last_project
from app.file_memory import remember_file, get_remembered_file
from app.app_finder import find_app
from app.software_indexer import find_software
from app.file_indexer import find_file
from pathlib import Path
import subprocess
import shutil
from datetime import datetime
from app.app_finder import find_app
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from app.software_indexer import find_software
from app.smart_search import find_similar_items
from app.session_memory import set_last_created_path, set_last_opened_path, resolve_reference
from app.path_resolver import resolve_location, is_blocked_path
from app.session_memory import (
    get_last_project,
    set_last_project,
    set_last_created_path,
    set_last_opened_path,
    set_last_active_path,
    resolve_reference
)

IGNORED_DIRS = {
    ".git", "venv", "__pycache__", "node_modules", ".idea", ".vscode",
    "dist", "build", ".env"
}

LAST_PROJECT_PATH = None
LAST_PROJECT_TEMPLATE = None


def set_last_project(path, template):
    global LAST_PROJECT_PATH, LAST_PROJECT_TEMPLATE

    LAST_PROJECT_PATH = str(path)
    LAST_PROJECT_TEMPLATE = template


def get_last_project():
    return LAST_PROJECT_PATH, LAST_PROJECT_TEMPLATE




def format_size(size_bytes: int):
    if size_bytes < 1024:
        return f"{size_bytes} B"
    if size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    if size_bytes < 1024 * 1024 * 1024:
        return f"{size_bytes / (1024 * 1024):.1f} MB"
    return f"{size_bytes / (1024 * 1024 * 1024):.1f} GB"


def build_file_item(path: Path):
    try:
        stat = path.stat()
        size = "" if path.is_dir() else format_size(stat.st_size)
        modified = datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %I:%M %p")
    except Exception:
        size = ""
        modified = ""

    return {
        "name": path.name,
        "type": "Folder" if path.is_dir() else "File",
        "size": size,
        "modified": modified,
        "path": str(path)
    }

def smart_open_file(query: str):
    query = query.strip()

    remembered = get_remembered_file(query)

    if remembered:
        remembered_path = Path(remembered)

        if remembered_path.exists():
            try:
                subprocess.Popen(f'explorer "{remembered_path}"', shell=True)
                set_last_opened_path(remembered_path)
                return True, f"Opened remembered file: {remembered_path}"
            except Exception:
                pass

    result = find_file(query)

    if result.get("found"):
        path = result["file"]["path"]
        remember_file(query, path)

        try:
            subprocess.Popen(f'explorer "{path}"', shell=True)
            set_last_opened_path(path)
            return True, f"Opened file: {path}"
        except Exception as e:
            return False, f"Failed to open file: {e}"

    suggestions = result.get("suggestions", [])

    if suggestions:
        return False, {
            "text": f"No exact match for '{query}'. Similar files found.",
            "viewer": {
                "title": f"Similar files for '{query}'",
                "items": [
                    {
                        "name": item.get("name", ""),
                        "type": "File",
                        "size": "",
                        "modified": "",
                        "score": "",
                        "path": item.get("path", "")
                    }
                    for item in suggestions
                ]
            }
        }

    return False, f"No files found for: {query}"

def should_ignore(path: Path):
    return any(part in IGNORED_DIRS for part in path.parts)

def refresh_desktop():
    try:
        subprocess.Popen("ie4uinit.exe -show", shell=True)
    except Exception:
        pass

def get_unique_path(path: Path):
    if not path.exists():
        return path

    counter = 1
    while True:
        new_path = path.parent / f"{path.name}_{counter}"
        if not new_path.exists():
            return new_path
        counter += 1


def get_unique_file_path(path: Path):
    if not path.exists():
        return path

    counter = 1
    while True:
        new_path = path.parent / f"{path.stem}_{counter}{path.suffix}"
        if not new_path.exists():
            return new_path
        counter += 1


def create_folder(folder_name: str, location: str):
    location = resolve_reference(location)
    folder_name = resolve_reference(folder_name)

    base_path = resolve_location(location)
    folder_path = base_path / folder_name

    if is_blocked_path(folder_path):
        return False, f"Blocked for safety: {folder_path}"

    try:
        folder_path = get_unique_path(folder_path)
        folder_path.mkdir(parents=True, exist_ok=False)

        set_last_created_path(folder_path)
        refresh_desktop()
        return True, {
            "text": f"Folder created successfully: {folder_path}",
            "created_path": str(folder_path)
        }

    except Exception as e:
        return False, f"Failed to create folder: {e}"


def create_file(file_name: str, location: str):
    location = resolve_reference(location)
    file_name = resolve_reference(file_name)

    base_path = resolve_location(location)
    file_path = base_path / file_name

    if is_blocked_path(file_path):
        return False, f"Blocked for safety: {file_path}"

    try:
        base_path.mkdir(parents=True, exist_ok=True)
        file_path = get_unique_file_path(file_path)

        suffix = file_path.suffix.lower()

        if suffix == ".docx":
            doc = Document()
            doc.add_heading("Created by Orvix", level=1)
            doc.add_paragraph("This document was generated by Orvix.")
            doc.save(file_path)

        elif suffix == ".xlsx":
            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            ws["A1"] = "Created by Orvix"
            wb.save(file_path)

        elif suffix == ".pptx":
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = "Created by Orvix"
            slide.placeholders[1].text = "This presentation was generated by Orvix."
            prs.save(file_path)

        else:
            file_path.write_text("", encoding="utf-8")

        set_last_created_path(file_path)
        refresh_desktop()
        return True, {
            "text": f"File created successfully: {file_path}",
            "created_path": str(file_path)
        }

    except Exception as e:
        return False, f"Failed to create file: {e}"


def open_folder(folder_name: str, location: str):
    location = resolve_reference(location)
    folder_name = resolve_reference(folder_name)

    base_path = resolve_location(location)

    if folder_name:
        folder_path = base_path / folder_name
    else:
        folder_path = base_path

    if is_blocked_path(folder_path):
        return False, f"Blocked for safety: {folder_path}"

    if folder_path.exists() and folder_path.is_dir():
        try:
            subprocess.Popen(f'explorer "{folder_path}"', shell=True)
            set_last_opened_path(folder_path)
            return True, f"Opened folder: {folder_path}"
        except Exception as e:
            return False, f"Failed to open folder: {e}"

    if not base_path.exists():
        return False, f"Location not found: {base_path}"

    similar = find_similar_items(base_path, folder_name, item_type="folder", limit=20)

    if similar:
        items = []
        for result in similar:
            item = build_file_item(result["path"])
            item["score"] = str(result["score"])
            items.append(item)

        return False, {
            "text": f"Folder not found exactly: {folder_path}. Showing similar folders.",
            "viewer": {
                "title": f"Similar folders for '{folder_name}'",
                "items": items
            }
        }

    return False, f"Folder not found: {folder_path}"


def open_file(file_name: str, location: str):
    location = resolve_reference(location)
    file_name = resolve_reference(file_name)

    base_path = resolve_location(location)
    file_path = base_path / file_name

    if is_blocked_path(file_path):
        return False, f"Blocked for safety: {file_path}"

    if not base_path.exists():
        return False, f"Location not found: {base_path}"

    if file_path.exists() and file_path.is_file():
        try:
            subprocess.Popen(f'explorer "{file_path}"', shell=True)
            set_last_opened_path(file_path)
            return True, f"Opened file: {file_path}"
        except Exception as e:
            return False, f"Failed to open file: {e}"

    stem_matches = []
    for item in base_path.iterdir():
        if item.is_file() and item.stem.lower() == file_name.lower():
            stem_matches.append(item)

    if len(stem_matches) == 1:
        matched_file = stem_matches[0]
        try:
            subprocess.Popen(f'explorer "{matched_file}"', shell=True)
            set_last_opened_path(matched_file)
            return True, f"Opened closest match: {matched_file}"
        except Exception as e:
            return False, f"Failed to open matched file: {e}"

    similar = find_similar_items(base_path, file_name, item_type="file", limit=20)

    if similar:
        items = []
        for result in similar:
            item = build_file_item(result["path"])
            item["score"] = str(result["score"])
            items.append(item)

        return False, {
            "text": f"File not found exactly: {file_path}. Showing similar files.",
            "viewer": {
                "title": f"Similar files for '{file_name}'",
                "items": items
            }
        }

    return False, f"File not found: {file_path}"


def open_app(app_name: str):
    app_name = app_name.lower().strip()

    built_in = {
        "vscode": "code",
        "vs code": "code",
        "visual studio code": "code",
        "chrome": "chrome",
        "google chrome": "chrome",
        "notepad": "notepad",
        "calculator": "calc",
        "paint": "mspaint",
        "cmd": "cmd",
        "terminal": "wt",
        "file explorer": "explorer",
        "word": "winword",
        "excel": "excel",
        "powerpoint": "powerpnt"
    }

    # 1. Built-in known commands
    if app_name in built_in:
        try:
            subprocess.Popen(built_in[app_name], shell=True)
            return True, f"Opened app: {app_name}"
        except Exception:
            pass

    # 2. Search Start Menu / Desktop shortcuts
    result = find_app(app_name)

    if result.get("found"):
        try:
            subprocess.Popen(f'explorer "{result["path"]}"', shell=True)
            return True, f"Opened app shortcut: {result.get('name', app_name)}"
        except Exception as e:
            return False, f"Failed to open app shortcut: {e}"

    # 3. Search Program Files / Program Files (x86)
    software_path = find_software(app_name)

    if software_path:
        try:
            subprocess.Popen(str(software_path), shell=True)
            return True, f"Opened software: {software_path.name}"
        except Exception as e:
            return False, f"Failed to open software: {e}"

    # 4. Show suggestions from shortcut index
    suggestions = result.get("suggestions", [])

    suggestions = result.get("suggestions", [])

    if suggestions:
        items = []

        for suggestion in suggestions:
            if isinstance(suggestion, dict):
                items.append({
                    "name": suggestion.get("name", ""),
                    "type": "Application",
                    "size": "",
                    "modified": "",
                    "score": "",
                    "path": suggestion.get("path", "")
                })
            else:
                items.append({
                    "name": str(suggestion),
                    "type": "Application",
                    "size": "",
                    "modified": "",
                    "score": "",
                    "path": ""
                })

        return False, {
            "text": f"App '{app_name}' not found exactly. Similar apps found. Double-click a row to open.",
            "viewer": {
                "title": "Similar Apps",
                "items": items
            }
        }
def open_app_in_location(app_name: str, location: str):
    location = resolve_reference(location)
    folder_path = resolve_location(location)

    if is_blocked_path(folder_path):
        return False, f"Blocked for safety: {folder_path}"

    if not folder_path.exists():
        return False, f"Folder not found: {folder_path}"

    app_name = app_name.lower().strip()

    if app_name in ["vscode", "vs code", "visual studio code"]:
        command = f'code "{folder_path}"'
    else:
        return False, f"Opening {app_name} in folder is not supported yet."

    try:
        subprocess.Popen(command, shell=True)
        set_last_opened_path(folder_path)
        return True, f"Opened {folder_path} in VS Code"
    except Exception as e:
        return False, f"Failed to open VS Code in folder: {e}"


def run_project(template: str, project_name: str, location: str):
    project_name = "" if project_name is None else str(project_name).strip()

    if project_name.lower() in ["", "it", "this", "that"]:
        last_path, last_template = get_last_project()

        if not last_path:
            return False, "No recent project found to run."

        project_path = Path(last_path)
        template = last_template or template

    else:
        location = resolve_reference(location)
        project_name = resolve_reference(project_name)

        base_path = resolve_location(location)
        project_path = base_path / project_name

    if is_blocked_path(project_path):
        return False, f"Blocked for safety: {project_path}"

    if not project_path.exists():
        return False, f"Project folder not found: {project_path}"

    commands = {
        "react app": "npm install && npm run dev",
        "node express app": "npm install && npm run dev",
        "django app": "pip install -r requirements.txt && python manage.py runserver",
        "flask app": "pip install -r requirements.txt && python run.py",
        "python app": "python app/main.py",
        "ml project": "pip install -r requirements.txt && python src/main.py",
        "data science project": "pip install -r requirements.txt",
        "spring boot project": "mvn spring-boot:run",
    }

    run_command = commands.get(template)

    if not run_command:
        return False, f"No run command found for template: {template}"

    try:
        subprocess.Popen(
            f'start cmd /k "cd /d {project_path} && {run_command}"',
            shell=True
        )
        set_last_opened_path(project_path)
        return True, f"Running {template}: {project_path}"
    except Exception as e:
        return False, f"Failed to run project: {e}"
    if is_blocked_path(project_path):
        return False, f"Blocked for safety: {project_path}"

    if not project_path.exists():
        return False, f"Project folder not found: {project_path}"

    commands = {
        "react app": "npm install && npm run dev",
        "node express app": "npm install && npm run dev",
        "django app": "pip install -r requirements.txt && python manage.py runserver",
        "flask app": "pip install -r requirements.txt && python run.py",
        "python app": "python app/main.py",
        "ml project": "pip install -r requirements.txt && python src/main.py",
        "data science project": "pip install -r requirements.txt",
        "spring boot project": "mvn spring-boot:run",
    }

    run_command = commands.get(template)

    if not run_command:
        return False, f"No run command found for template: {template}"

    try:
        subprocess.Popen(
            f'start cmd /k "cd /d {project_path} && {run_command}"',
            shell=True
        )
        set_last_opened_path(project_path)
        return True, f"Running {template}: {project_path}"
    except Exception as e:
        return False, f"Failed to run project: {e}"


def copy_or_move_item(operation: str, item_type: str, name: str, source: str, destination: str):
    source = resolve_reference(source)
    destination = resolve_reference(destination)
    name = resolve_reference(name)

    source_base = resolve_location(source)
    destination_base = resolve_location(destination)

    source_path = source_base / name
    destination_path = destination_base / name

    if is_blocked_path(source_path) or is_blocked_path(destination_path):
        return False, "Blocked for safety: system path detected."

    if not source_path.exists():
        return False, f"Source not found: {source_path}"

    if destination_path.exists():
        return False, f"Destination already exists. Overwrite blocked: {destination_path}"

    destination_base.mkdir(parents=True, exist_ok=True)

    try:
        if operation == "copy":
            if item_type == "file":
                if not source_path.is_file():
                    return False, f"Source is not a file: {source_path}"
                shutil.copy2(source_path, destination_path)

            elif item_type == "folder":
                if not source_path.is_dir():
                    return False, f"Source is not a folder: {source_path}"
                shutil.copytree(source_path, destination_path)

            else:
                return False, f"Unsupported item type: {item_type}"

            set_last_created_path(destination_path)
            refresh_desktop()
            return True, f"Copied {item_type}: {source_path} → {destination_path}"

        if operation == "move":
            if item_type == "file" and not source_path.is_file():
                return False, f"Source is not a file: {source_path}"

            if item_type == "folder" and not source_path.is_dir():
                return False, f"Source is not a folder: {source_path}"

            shutil.move(str(source_path), str(destination_path))
            set_last_created_path(destination_path)
            return True, f"Moved {item_type}: {source_path} → {destination_path}"

        return False, f"Unsupported operation: {operation}"

    except Exception as e:
        return False, f"Failed to {operation}: {e}"


def rename_item(item_type: str, old_name: str, new_name: str, location: str):
    location = resolve_reference(location)
    old_name = resolve_reference(old_name)

    base_path = resolve_location(location)

    old_path = base_path / old_name
    new_path = base_path / new_name

    if is_blocked_path(old_path) or is_blocked_path(new_path):
        return False, "Blocked for safety: system path detected."

    if not old_path.exists():
        return False, f"Source not found: {old_path}"

    if new_path.exists():
        return False, f"Target already exists. Rename blocked: {new_path}"

    if item_type == "file" and not old_path.is_file():
        return False, f"Source is not a file: {old_path}"

    if item_type == "folder" and not old_path.is_dir():
        return False, f"Source is not a folder: {old_path}"

    try:
        old_path.rename(new_path)
        set_last_created_path(new_path)
        refresh_desktop()
        return True, f"Renamed {item_type}: {old_path} → {new_path}"
    except Exception as e:
        return False, f"Failed to rename: {e}"


def delete_item(item_type: str, name: str, location: str):
    location = resolve_reference(location)
    name = resolve_reference(name)

    base_path = resolve_location(location)
    target_path = base_path / name

    if is_blocked_path(target_path):
        return False, f"Blocked for safety: {target_path}"

    if not target_path.exists():
        return False, f"Target not found: {target_path}"

    if item_type == "file" and not target_path.is_file():
        return False, f"Target is not a file: {target_path}"

    if item_type == "folder" and not target_path.is_dir():
        return False, f"Target is not a folder: {target_path}"

    try:
        if item_type == "file":
            target_path.unlink()
        else:
            shutil.rmtree(target_path)
        refresh_desktop()
        return True, f"Deleted {item_type}: {target_path}"

    except Exception as e:
        return False, f"Failed to delete: {e}"


def list_folder(location: str):
    location = resolve_reference(location)
    folder_path = resolve_location(location)

    if is_blocked_path(folder_path):
        return False, f"Blocked for safety: {folder_path}"

    if not folder_path.exists():
        return False, f"Folder not found: {folder_path}"

    if not folder_path.is_dir():
        return False, f"Target is not a folder: {folder_path}"

    try:
        items = []

        for item in folder_path.iterdir():
            if should_ignore(item):
                continue
            items.append(build_file_item(item))

        set_last_opened_path(folder_path)

        return True, {
            "text": f"Listed {len(items)} items from {folder_path}",
            "viewer": {
                "title": f"Contents of {folder_path}",
                "items": items
            }
        }

    except Exception as e:
        return False, f"Failed to list folder: {e}"


def search_file(query: str, location: str):
    location = resolve_reference(location)
    folder_path = resolve_location(location)

    if is_blocked_path(folder_path):
        return False, f"Blocked for safety: {folder_path}"

    if not folder_path.exists():
        return False, f"Search location not found: {folder_path}"

    if not folder_path.is_dir():
        return False, f"Search location is not a folder: {folder_path}"

    query_lower = query.lower()
    matches = []

    try:
        for path in folder_path.rglob("*"):
            if should_ignore(path):
                continue

            if is_blocked_path(path):
                continue

            if query_lower in path.name.lower():
                matches.append(build_file_item(path))

            if len(matches) >= 100:
                break

        if not matches:
            return False, f"No matching files found for '{query}' in {folder_path}"

        set_last_opened_path(folder_path)

        return True, {
            "text": f"Found {len(matches)} results for '{query}' in {folder_path}",
            "viewer": {
                "title": f"Search results for '{query}'",
                "items": matches
            }
        }

    except Exception as e:
        return False, f"Search failed: {e}"
    
def create_project_structure(project_name: str, location: str, folders: list, files, template="react app"):
    location = resolve_reference(location)
    project_name = resolve_reference(project_name)

    base_path = resolve_location(location)
    project_path = base_path / project_name

    if is_blocked_path(project_path):
        return False, f"Blocked for safety: {project_path}"

    original_project_path = project_path
    project_path = get_unique_path(project_path)

    if project_path != original_project_path:
        renamed_message = f"Project already existed. Created unique project: {project_path}"
    else:
        renamed_message = f"Project created: {project_path}"

    try:
        project_path.mkdir(parents=True, exist_ok=False)
        messages = [renamed_message]

        for folder in folders:
            folder_path = project_path / folder

            if is_blocked_path(folder_path):
                messages.append(f"Skipped blocked folder: {folder_path}")
                continue

            folder_path.mkdir(parents=True, exist_ok=True)
            messages.append(f"Folder created: {folder_path}")

        if isinstance(files, dict):
            file_items = files.items()
        else:
            file_items = [(file, None) for file in files]

        for file_name, content in file_items:
            file_path = project_path / file_name

            if is_blocked_path(file_path):
                messages.append(f"Skipped blocked file: {file_path}")
                continue

            file_path.parent.mkdir(parents=True, exist_ok=True)

            if content is None:
                success, msg = create_file(file_path.name, str(file_path.parent))

                if isinstance(msg, dict):
                    messages.append(msg.get("text", ""))
                else:
                    messages.append(msg)
            else:
                file_path.write_text(content, encoding="utf-8")
                messages.append(f"File created with starter code: {file_path}")

                set_last_created_path(project_path)
                set_last_project(project_path, template)
                refresh_desktop()
                
        return True, {
            "text": "\n".join(messages),
            "created_path": str(project_path)
        }

    except Exception as e:
        return False, f"Failed to create project: {e}"


def execute_single_action(parsed_command):
    action = parsed_command.get("action")

    if action == "create_ai_word_document":
        from app.word_command_handler import create_word_document_from_command

        result = create_word_document_from_command(
            parsed_command.get("command", "")
        )

        return (
            result.get("success", False),
            result.get("message", "Word document task finished.")
        )

    if action == "document_action":
        from app.word_agent.document_actions import execute_document_action

        result = execute_document_action(
            parsed_command.get("command", "")
        )

        return (
            result.get("success", False),
            result.get("message", "Document action finished.")
        )

    if action == "modify_current_project":
        from app.intent_planner import handle_modify_current_project

        result = handle_modify_current_project(
            parsed_command.get("command", "")
        )

        if isinstance(result, dict):
            return (
                result.get("success", False),
                result.get("message", "Project modification finished.")
            )

        return True, result

    if action == "run_project":
        return run_project(
            parsed_command.get("template"),
            parsed_command.get("name"),
            parsed_command.get("location")
        )

    if action == "create_project":
     return create_project_structure(
        parsed_command.get("name"),
        parsed_command.get("location"),
        parsed_command.get("folders", []),
        parsed_command.get("files", []),
        parsed_command.get("template", "react app")
    )

    if action == "open_app_in_location":
        return open_app_in_location(
            parsed_command.get("app"),
            parsed_command.get("location")
        )

    if action == "create_folder":
        return create_folder(parsed_command.get("name"), parsed_command.get("location"))

    if action == "create_file":
        return create_file(parsed_command.get("name"), parsed_command.get("location"))

    if action == "open_folder":
        return open_folder(parsed_command.get("name"), parsed_command.get("location"))

    if action == "open_file":
        return open_file(parsed_command.get("name"), parsed_command.get("location"))

    if action == "open_app":
        return open_app(parsed_command.get("app"))

    if action in ["copy", "move"]:
        return copy_or_move_item(
            action,
            parsed_command.get("item_type"),
            parsed_command.get("name"),
            parsed_command.get("source"),
            parsed_command.get("destination")
        )

    if action == "rename":
        return rename_item(
            parsed_command.get("item_type"),
            parsed_command.get("old_name"),
            parsed_command.get("new_name"),
            parsed_command.get("location")
        )

    if action == "delete":
        return delete_item(
            parsed_command.get("item_type"),
            parsed_command.get("name"),
            parsed_command.get("location")
        )

    if action == "list_folder":
        return list_folder(parsed_command.get("location"))

    if action == "search_file":
        return search_file(parsed_command.get("query"), parsed_command.get("location"))

    if action == "smart_open_file":
        return smart_open_file(parsed_command.get("query"))

    return False, "No executor found for this action."


def execute_action(parsed_command):
    if parsed_command.get("action") == "multi_step":
        messages = []
        final_success = True
        viewer_payload = None
        last_created_path = None

        for step in parsed_command.get("steps", []):
            if step.get("use_last_created_path") and last_created_path:
                step["location"] = last_created_path

            success, message = execute_single_action(step)

            if isinstance(message, dict):
                messages.append(("✓ " if success else "✗ ") + message.get("text", ""))

                if message.get("viewer"):
                    viewer_payload = message.get("viewer")

                if message.get("created_path"):
                    last_created_path = message.get("created_path")
            else:
                messages.append(("✓ " if success else "✗ ") + message)

            if not success:
                final_success = False
                break

        if viewer_payload:
            return final_success, {
                "text": "\n".join(messages),
                "viewer": viewer_payload
            }

        return final_success, "\n".join(messages)

    return execute_single_action(parsed_command)