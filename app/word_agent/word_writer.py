from docx import Document


def append_to_word_document(file_path, heading, content):
    document = Document(file_path)

    if heading:
        document.add_heading(heading, level=1)

    for block in content.split("\n"):
        block = block.strip()

        if block:
            document.add_paragraph(block)

    document.save(file_path)

    return {
        "success": True,
        "path": file_path,
        "message": f"Updated Word document: {file_path}"
    }