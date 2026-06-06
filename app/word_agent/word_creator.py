import os

from docx import Document
from app.word_agent.word_memory import save_current_document


def get_desktop_path():
    onedrive_desktop = os.path.join(
        os.path.expanduser("~"),
        "OneDrive",
        "Desktop"
    )

    if os.path.exists(onedrive_desktop):
        return onedrive_desktop

    return os.path.join(os.path.expanduser("~"), "Desktop")


def get_unique_file_path(folder, filename):
    base, ext = os.path.splitext(filename)
    file_path = os.path.join(folder, filename)

    counter = 1

    while os.path.exists(file_path):
        file_path = os.path.join(folder, f"{base}_{counter}{ext}")
        counter += 1

    return file_path


def create_word_document(name, title=None, content=None, open_file=True):
    desktop = get_desktop_path()

    if not name.lower().endswith(".docx"):
        name = f"{name}.docx"

    file_path = get_unique_file_path(desktop, name)

    document = Document()

    if title:
        document.add_heading(title, 0)

    if content:
        document.add_paragraph(content)

    document.save(file_path)
    save_current_document(file_path)

    if open_file:
        os.startfile(file_path)

    return {
        "success": True,
        "path": file_path,
        "message": f"Word document created: {file_path}"
    }