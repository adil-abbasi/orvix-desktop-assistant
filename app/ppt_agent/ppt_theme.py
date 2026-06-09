from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


THEMES = {
    "default": {
        "background": RGBColor(15, 23, 42),
        "primary": RGBColor(96, 165, 250),
        "secondary": RGBColor(129, 140, 248),
        "text": RGBColor(248, 250, 252),
        "muted": RGBColor(203, 213, 225),
        "card": RGBColor(30, 41, 59),
    },
    "ai": {
        "background": RGBColor(2, 6, 23),
        "primary": RGBColor(34, 211, 238),
        "secondary": RGBColor(168, 85, 247),
        "text": RGBColor(248, 250, 252),
        "muted": RGBColor(186, 230, 253),
        "card": RGBColor(15, 23, 42),
    },
    "business": {
        "background": RGBColor(17, 24, 39),
        "primary": RGBColor(59, 130, 246),
        "secondary": RGBColor(16, 185, 129),
        "text": RGBColor(249, 250, 251),
        "muted": RGBColor(209, 213, 219),
        "card": RGBColor(31, 41, 55),
    },
    "medical": {
        "background": RGBColor(236, 253, 245),
        "primary": RGBColor(5, 150, 105),
        "secondary": RGBColor(14, 165, 233),
        "text": RGBColor(6, 78, 59),
        "muted": RGBColor(55, 65, 81),
        "card": RGBColor(255, 255, 255),
    },
}


def detect_theme(topic: str):
    text = topic.lower()

    if any(word in text for word in ["ai", "artificial intelligence", "machine learning", "data science"]):
        return "ai"

    if any(word in text for word in ["business", "startup", "company", "marketing"]):
        return "business"

    if any(word in text for word in ["health", "medical", "doctor", "hospital", "disease"]):
        return "medical"

    return "default"


def get_theme(topic: str):
    theme_name = detect_theme(topic)
    return THEMES.get(theme_name, THEMES["default"])


def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, text, theme, top=0.7, left=0.65, width=12.0, height=1.0, size=36):
    box = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height)
    )

    frame = box.text_frame
    frame.clear()

    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.bold = True
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = theme["text"]
    paragraph.alignment = PP_ALIGN.LEFT

    return box


def add_subtitle(slide, text, theme, top=1.65, left=0.7, width=11.5, height=0.8, size=18):
    box = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height)
    )

    frame = box.text_frame
    frame.clear()

    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = theme["muted"]

    return box


def add_footer(slide, text, theme, slide_number=None):
    footer_text = text

    if slide_number is not None:
        footer_text = f"{text}  •  Slide {slide_number}"

    box = slide.shapes.add_textbox(
        Inches(0.65),
        Inches(7.05),
        Inches(12.0),
        Inches(0.3)
    )

    frame = box.text_frame
    frame.clear()

    paragraph = frame.paragraphs[0]
    paragraph.text = footer_text
    paragraph.font.size = Pt(9)
    paragraph.font.color.rgb = theme["muted"]

    return box


def add_accent_bar(slide, theme):
    shape = slide.shapes.add_shape(
        1,
        Inches(0),
        Inches(0),
        Inches(13.33),
        Inches(0.12)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = theme["primary"]
    shape.line.fill.background()


def add_card(slide, left, top, width, height, theme):
    card = slide.shapes.add_shape(
        5,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = theme["card"]
    card.line.color.rgb = theme["secondary"]
    card.line.width = Pt(1)

    return card